"""Audit a local MAT collection without changing the operational database or contacting SIH."""
from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

import numpy as np
import pandas as pd
import yaml
from openpyxl import load_workbook

from aforix.analysis.uncertainty.calculator import IVEParameters
from aforix.export.sih.ive_integration import prepare_sih_ive
from aforix.export.sih.schema import SDH_ACTUACIONES_COLUMNS, SDH_AFOROS_COLUMNS
from aforix.ingest.adapters.m9_mat import parse_m9_mat
from aforix.validation.m9 import validate_m9

KEYS = ['instrument', 'station_id', 'measurement_date', 'measurement_time']


def check_export(root, measurement, metadata, parameters, sih, reviews=None):
    """Check serialized CSV/XLSX against normalized data and the shared engine."""
    errors = []
    if reviews is None:
        reviews = []
    def check(condition, message):
        if not condition:
            errors.append(message)
    def equal(actual, expected, label):
        try:
            check(bool(np.isclose(float(actual), float(expected), rtol=1e-8, atol=1e-10)), label)
        except (ValueError, TypeError):
            errors.append(label)
    output = sih['sih'].get('output', {})
    csv_args = dict(dtype=str, keep_default_na=False,
                    sep=output.get('delimiter', ','), encoding=output.get('encoding', 'utf-8-sig'))
    aforos = pd.read_csv(metadata['aforos_file'], **csv_args)
    actuaciones = pd.read_csv(metadata['actuaciones_file'], **csv_args)
    check(list(aforos.columns) == list(SDH_AFOROS_COLUMNS), 'Esquema CSV aforo')
    check(list(actuaciones.columns) == list(SDH_ACTUACIONES_COLUMNS), 'Esquema CSV actuacion')
    check(len(aforos) == len(actuaciones) == 1, 'Cantidad de filas SIH')
    if len(aforos) != 1 or len(actuaciones) != 1:
        return errors
    aforo, actuacion = aforos.iloc[0], actuaciones.iloc[0]
    for target, source in [('caudal', 'q_total_m3s'), ('seccion', 'area_total_m2'),
                           ('ancho', 'width_total_m'), ('profundidad', 'depth_mean_m'),
                           ('velocidad_media', 'velocity_mean_m_s'), ('radio_hidraulico', 'hydraulic_radius_m')]:
        expected = measurement[source]
        if pd.isna(expected) or str(expected).strip() == '':
            reviews.append(f'{source}: dato normalizado ausente')
            check(aforo[target] == '', f'{target}: valor exportado sin dato normalizado')
        else:
            equal(aforo[target], expected, target)
    check(aforo['id_estacion'] == str(measurement['station_id']), 'Estacion aforo')
    check(actuacion['id_estacion'] == aforo['id_estacion'], 'Estacion actuacion')
    catalog = pd.read_csv(root/'configs/sih/instrumentos.csv', dtype=str, keep_default_na=False)
    found = catalog[(catalog['marca'] == measurement['manufacturer']) &
                    (catalog['nro_serie'] == measurement['serial_number']) &
                    (catalog['modelo'] == measurement['model'])]
    check(len(found) == 1 and aforo['id_instrumento'] == found.iloc[0]['id'], 'Instrumento SIH/catalogo')
    fmt = sih['sih'].get('datetime', {}).get('output_format', '%d/%m/%Y %H:%M:%S')
    for target, datecol, timecol in [('fecha_inicio', 'measurement_date', 'measurement_time'),
                                     ('fecha_fin', 'measurement_end_date', 'measurement_end_time')]:
        expected = datetime.strptime(measurement[datecol] + measurement[timecol], '%Y%m%d%H%M%S').strftime(fmt)
        check(aforo[target] == expected, target)
    try:
        selected, result = prepare_sih_ive(root/'database/normalized', 'm9', measurement, parameters)
    except (ValueError, FileNotFoundError, KeyError, TypeError, OverflowError):
        check(metadata['ive_status'] == 'not_available', 'Estado IVE no disponible')
        check(aforo['error_IVE'] == aforo['incertidumbre_IVE'] == '', 'IVE no disponible debe estar vacio')
        return errors
    check(metadata['ive_status'] == result.status, 'Estado IVE')
    if result.status == 'calculated':
        equal(aforo['error_IVE'], result.absolute_uncertainty_m3s, 'Error IVE')
        equal(aforo['incertidumbre_IVE'], result.relative_uncertainty_pct, 'Porcentaje IVE')
        check(aforo['calidad_IVE'] == result.quality, 'Calidad IVE')
    else:
        check(aforo['error_IVE'] == aforo['incertidumbre_IVE'] == '', 'IVE no calculado debe estar vacio')
    report = str(metadata.get('ive_report_file', ''))
    check(bool(report) and Path(report).is_file(), 'Informe XLSX ausente')
    if report and Path(report).is_file():
        workbook = load_workbook(report, data_only=True, read_only=True)
        try:
            values = dict(workbook['Resumen'].iter_rows(min_row=2, max_col=2, values_only=True))
            check(values['Estado IVE'] == result.status, 'Estado XLSX')
            equal(values['Caudal Mid-section (m3/s)'], result.q_mid_m3s, 'Q_mid XLSX')
            check(workbook['Points_normalizados'].max_row - 1 == len(selected), 'Muestras XLSX')
            if result.status == 'calculated':
                equal(values['Incertidumbre absoluta (m3/s)'], aforo['error_IVE'], 'IVE absoluto XLSX/CSV')
                equal(values['Incertidumbre relativa (%)'], aforo['incertidumbre_IVE'], 'IVE relativo XLSX/CSV')
        finally:
            workbook.close()
    return errors


def isolate_paths(cfg, sih):
    """Force all pipeline writes and data inputs into the audit project."""
    cfg['paths'].update(raw_data_dir='data/raw', runs_root='runs', database_root='database')
    cfg['ingest']['m9'].update(enabled=True, raw_subdir='M9')
    cfg['build_groups'].update(enabled=True, sources=['m9'], input_runs_root='runs',
                               output_dir='database/raw_canonical', groups=['Summary', 'Points'],
                               concat_groups=['Summary'], use_latest_run_only=True)
    cfg['normalize'].update(enabled=True, sources=['m9'], input_dir='database/raw_canonical',
                            output_dir='database/normalized', registry_dir='configs/normalization',
                            groups=['Summary', 'Points'], concat_groups=['Summary', 'Points'],
                            write_policy='overwrite')
    cfg['validation'].update(enabled=True, input_dir='database/normalized', output_dir='database/validation')
    sih['sih']['inputs'].update(normalized_input_dir='database/normalized',
                                raw_canonical_input_dir='database/raw_canonical',
                                quality_input_dir='database/analysis/quality_metrics')
    sih['sih']['output'].update(output_dir='outputs/sih', output_names={
        'actuaciones': 'ID_{export_id}_actuacion_{station_id}_{YYYYMMDD}_{HHMMSS}.csv',
        'aforos': 'ID_{export_id}_aforo_{station_id}_{YYYYMMDD}_{HHMMSS}.csv',
    })


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--raw-dir', type=Path, required=True)
    parser.add_argument('--output-root', type=Path, default=Path('runs/audit_m9'))
    parser.add_argument('--exclude', action='append', default=[], help='Exact filename or relative path; repeatable')
    parser.add_argument('--pattern', default='*.mat', help='Recursive filename pattern')
    args = parser.parse_args(argv)
    raw_source = args.raw_dir.resolve()
    if not raw_source.is_dir():
        parser.error(f'No existe la carpeta: {raw_source}')
    files = sorted(p for p in raw_source.rglob(args.pattern) if p.is_file() and p.suffix.lower() == '.mat')
    if not files:
        parser.error('No se encontraron MAT; no se ejecuto el proceso.')
    repo = Path(__file__).resolve().parents[1]
    args.output_root.mkdir(parents=True, exist_ok=True)
    root = Path(tempfile.mkdtemp(prefix=datetime.now().strftime('%Y%m%d_%H%M%S_'), dir=args.output_root.resolve()))
    # config_root_from_path searches repository markers before configs/.
    # Establish the isolated project boundary even when nested in the real repo.
    (root/'pyproject.toml').write_text('[tool.aforix_audit]\nisolated = true\n', encoding='utf-8')
    from aforix.config.paths import config_root_from_path
    if config_root_from_path(root/'configs/examples/main.yaml') != root:
        raise RuntimeError('No se pudo establecer una raiz de auditoria aislada')
    print(f'AUDIT_ROOT {root}', flush=True)
    shutil.copytree(repo/'configs', root/'configs')
    (root/'runs').mkdir()
    raw = root/'data/raw/M9'; raw.mkdir(parents=True)
    records = []
    mapping = {}
    for source in files:
        relative = source.relative_to(raw_source)
        excluded = source.name in args.exclude or relative.as_posix() in args.exclude
        record = dict(source_file=str(source), sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                      status='EXCLUIDO' if excluded else 'PENDIENTE', ingest='', validation='',
                      export='', ive='', message='Exclusion explicita' if excluded else '')
        records.append(record)
        if excluded:
            continue
        target = raw / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
        mapping[str(target.resolve())] = record
    manifest = root/'audit_manifest.csv'
    def save():
        pd.DataFrame(records).to_csv(manifest, index=False, encoding='utf-8-sig')
    save()
    cfgpath = root/'configs/examples/main.yaml'
    cfg = yaml.safe_load(cfgpath.read_text(encoding='utf-8'))
    sihpath = root/'configs/sih/sih.yaml'
    sih = yaml.safe_load(sihpath.read_text(encoding='utf-8'))
    isolate_paths(cfg, sih)
    # Snapshot external lookup files too, retaining the user's configured catalogs.
    for name, value in sih['sih'].get('lookup_files', {}).items():
        source = Path(value)
        source = source if source.is_absolute() else repo / source
        target = root / 'configs/sih' / f'{name}.csv'
        if source.is_file():
            shutil.copy2(source, target)
        sih['sih']['lookup_files'][name] = f'configs/sih/{name}.csv'
    cfgpath.write_text(yaml.safe_dump(cfg, sort_keys=False), encoding='utf-8')
    sihpath.write_text(yaml.safe_dump(sih, sort_keys=False), encoding='utf-8')
    borders = sih['sih'].get('ive', {}).get('borders', {})
    parameters = IVEParameters(left_bank_mode=borders.get('left', 'zero'), right_bank_mode=borders.get('right', 'preserve_depth'))
    def command(name, *arguments):
        print(f'Etapa: {name}', flush=True)
        with (root/f'{name}.log').open('w', encoding='utf-8') as log:
            result = subprocess.run([sys.executable, '-c', 'from aforix.cli.main import app; app()',
                                     *arguments, '-c', str(cfgpath)], cwd=root, stdout=log, stderr=subprocess.STDOUT)
        return result.returncode
    try:
        if not mapping:
            raise RuntimeError('Todos los archivos fueron excluidos; no hay aforos para auditar.')
        ingest_code = command('01_ingest', 'ingest', 'm9')
        statuses = list((root/'runs/ingest_m9').glob('*/ingest_status.json'))
        if len(statuses) != 1:
            log_path = root/'01_ingest.log'
            tail = log_path.read_text(encoding='utf-8', errors='replace')[-3000:]
            print(tail, flush=True)
            raise RuntimeError(f'No hay manifiesto de ingesta aislado (exit={ingest_code}); ver {log_path}')
        status = json.loads(statuses[0].read_text(encoding='utf-8'))
        for failed in status['failed']:
            rec = mapping[failed['file']]
            rec.update(status='ERROR', ingest='error', message=failed['error'])
        if ingest_code and not status['failed']:
            raise RuntimeError('Fallo de ingesta no explicado por el manifiesto')
        save()
        if not status['succeeded']:
            raise RuntimeError('Ningun archivo pudo ser ingerido')
        for name, arguments in [('02_groups', ['build-groups']), ('03_normalize', ['normalize', 'run']),
                                ('04_validate', ['validate', 'run'])]:
            if command(name, *arguments):
                raise RuntimeError(f'Fallo en {name}; revisar su log')
        normalized = root/'database/normalized/m9'
        summary = pd.read_csv(normalized/'Summary.csv', dtype=str, keep_default_na=False)
        points = pd.concat([pd.read_csv(p, dtype=str, keep_default_na=False) for p in (normalized/'Points').glob('*.csv')], ignore_index=True)
        duplicate = summary.duplicated(KEYS, keep=False)
        selection = []
        issues_all = []
        for index, row in summary.iterrows():
            rec = mapping[row['source_file']]
            rec['ingest'] = 'ok'
            source_points = points[points['source_file'] == row['source_file']]
            issues = validate_m9(summary.loc[[index]], source_points)
            if not issues.empty:
                issues.loc[issues['source_file'].fillna('').eq(''), 'source_file'] = row['source_file']
            issues_all.append(issues)
            rec['validation'] = 'ok' if issues.empty else 'revisar'
            rec['message'] = '; '.join(issues['message'].drop_duplicates().astype(str))
            original = parse_m9_mat(Path(row['source_file']))
            if len(source_points) != len(original['points']) or not np.isclose(float(row['q_total_m3s']), original['summary']['q_total_m3s']):
                rec.update(status='ERROR', message='Trazabilidad MAT/normalizado incorrecta')
                continue
            if duplicate.loc[index]:
                rec.update(status='REVISAR', export='bloqueada', message='Clave SIH ambigua: misma estacion/fecha/hora en distintos MAT; no se elige ni promedia automaticamente')
                continue
            export_id = f'M9_AUDIT_{index:04d}'
            rec['export_id'] = export_id
            rec['status'] = 'REVISAR' if not issues.empty else 'PENDIENTE'
            selection.append({**row[KEYS].to_dict(), 'export_id': export_id})
        pd.concat(issues_all, ignore_index=True).to_csv(root/'audit_issues.csv', index=False, encoding='utf-8-sig')
        save()
        if selection:
            selection_path = root/'selection.csv'
            pd.DataFrame(selection).to_csv(selection_path, index=False)
            if command('05_export', 'export', 'sih', '--sih-config', str(sihpath), '--selection-file', str(selection_path), '--yes'):
                raise RuntimeError('Fallo del comando export SIH; revisar 05_export.log')
            metadata = pd.read_csv(root/'outputs/sih/sih_export_metadata.csv', dtype=str, keep_default_na=False)
            print('Etapa: 06_check_exports', flush=True)
            for position, (_, meta) in enumerate(metadata.iterrows(), 1):
                print(f'Verificando exportacion {position}/{len(metadata)}', flush=True)
                rec = next(r for r in records if r.get('export_id') == meta['export_id'])
                rec['export'] = meta['status']; rec['ive'] = meta.get('ive_status', '')
                if meta['status'] != 'success':
                    rec.update(status='ERROR', message=meta.get('error', 'Exportacion fallida'))
                    continue
                row = summary.iloc[int(meta['export_id'].rsplit('_', 1)[-1])]
                try:
                    reviews = []
                    errors = check_export(root, row, meta, parameters, sih, reviews)
                except Exception as exc:
                    errors = [f'No se pudo verificar la salida: {exc}']
                rec['aforo_file'] = meta['aforos_file']; rec['actuacion_file'] = meta['actuaciones_file']
                rec['ive_report_file'] = meta.get('ive_report_file', '')
                if errors:
                    rec.update(status='ERROR', message='; '.join(errors))
                elif reviews or rec['validation'] != 'ok' or rec['ive'] != 'calculated':
                    reasons = [rec.get('message', ''), *reviews]
                    if rec['ive'] != 'calculated':
                        reasons.append(meta.get('ive_reason', ''))
                        reasons.append(row.get('ive_reason', ''))
                    rec.update(status='REVISAR', message='; '.join(dict.fromkeys(
                        reason for reason in reasons if reason)))
                else:
                    rec.update(status='OK', message='CSV/XLSX y trazabilidad verificados')
        for rec in records:
            if rec['status'] == 'PENDIENTE':
                rec.update(status='ERROR', message='Archivo sin resultado completo de auditoria')
    except Exception as exc:
        for rec in records:
            if rec['status'] == 'PENDIENTE':
                rec.update(status='ERROR', message=str(exc))
        (root/'fatal_error.txt').write_text(str(exc), encoding='utf-8')
        print(f'ERROR: {exc}', flush=True)
    save()
    counts = pd.Series([r['status'] for r in records]).value_counts().to_dict()
    report = {'counts': counts, 'files_found': len(files), 'raw_dir': str(raw_source),
              'actual_sih_import_tested': False, 'configuration': str(cfgpath)}
    (root/'audit_summary.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    (root/'LEEME.txt').write_text(
        'AUDITORIA LOCAL M9\n' + json.dumps(counts, ensure_ascii=False, indent=2) +
        '\nOK: controles locales aprobados. REVISAR: datos/IVE/seleccion requieren revision.\n'
        'ERROR: fallo de procesamiento o discrepancia. EXCLUIDO: no evaluado.\n'
        'audit_manifest.csv: resultado por MAT y rutas de salida. audit_issues.csv: incidencias.\n'
        'Los CSV generados son de prueba y no fueron importados al SIH.\n'
        'No se modificaron los originales ni la base operativa. No se promediaron transectos.\n', encoding='utf-8')
    print(json.dumps(counts, ensure_ascii=False), flush=True)
    print(f'Informe: {manifest}', flush=True)
    return 1 if counts.get('ERROR') or counts.get('REVISAR') or not counts.get('OK') else 0


if __name__ == '__main__':
    raise SystemExit(main())
