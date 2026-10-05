from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


def _norm_date(value: object) -> str:
    text = str(value).strip()
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    compact = "".join(ch for ch in text if ch.isdigit())[:8]
    if len(compact) == 8:
        return f"{compact[:4]}-{compact[4:6]}-{compact[6:8]}"
    return text


def _latest_workbook(root: Path) -> Path:
    candidates = sorted(
        root.glob("*/*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise SystemExit(f"No section-profile workbook found under {root}")
    return candidates[0]


def _read_measurement_data(workbook: Path, sheet: str) -> pd.DataFrame:
    data = pd.read_excel(workbook, sheet_name=sheet, skiprows=12)
    if "distance_m" in data.columns and "depth_m" in data.columns:
        return data

    raw = pd.read_excel(workbook, sheet_name=sheet, header=None)
    header_idx = None
    for i in range(len(raw)):
        vals = {str(v) for v in raw.iloc[i].dropna().tolist()}
        if {"distance_m", "depth_m"}.issubset(vals):
            header_idx = i
            break

    if header_idx is None:
        raise ValueError(f"{sheet}: data header not found")

    return pd.read_excel(workbook, sheet_name=sheet, skiprows=header_idx)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--station", required=True)
    parser.add_argument("--instrument", required=True)
    parser.add_argument("--date", required=True)
    parser.add_argument(
        "--root",
        default="runs_acceptance/analysis_section_profiles",
    )
    parser.add_argument(
        "--workbook",
        default=None,
        help="Explicit workbook path. If omitted, use the latest workbook under --root.",
    )
    args = parser.parse_args()

    root = Path(args.root)
    workbook = Path(args.workbook) if args.workbook else _latest_workbook(root)
    if not workbook.exists():
        raise SystemExit(f"Workbook not found: {workbook}")

    xls = pd.ExcelFile(workbook)
    checks: list[dict[str, str]] = []

    has_readme = "README" in xls.sheet_names
    has_index = "Index" in xls.sheet_names
    checks.append({
        "check": "workbook_structure",
        "status": "PASS" if has_readme and has_index else "FAIL",
        "detail": f"README={has_readme}, Index={has_index}, sheets={len(xls.sheet_names)}",
    })

    if not has_index:
        report = pd.DataFrame(checks)
        print(report.to_string(index=False))
        raise SystemExit(1)

    index = pd.read_excel(workbook, sheet_name="Index", dtype=str).fillna("")
    target_date = _norm_date(args.date)

    identity_ok = (
        not index.empty
        and set(index["station_id"]) == {str(args.station)}
        and set(index["instrument"]) == {str(args.instrument)}
        and set(index["measurement_date"].map(_norm_date)) == {target_date}
    )
    checks.append({
        "check": "authoritative_identity",
        "status": "PASS" if identity_ok else "FAIL",
        "detail": (
            f"rows={len(index)}, stations={sorted(set(index.get('station_id', [])))}, "
            f"instruments={sorted(set(index.get('instrument', [])))}, "
            f"dates={sorted(set(index.get('measurement_date', pd.Series(dtype=str)).map(_norm_date)))}"
        ),
    })

    no_p_alias = not index["station_id"].str.match(r"(?i)^p\d+$").any()
    checks.append({
        "check": "no_legacy_p_alias",
        "status": "PASS" if no_p_alias else "FAIL",
        "detail": f"station_ids={sorted(set(index['station_id']))}",
    })

    n_rows_ok = True
    n_rows_details: list[str] = []
    for _, row in index.iterrows():
        try:
            n_rows = int(float(row["n_rows"]))
        except Exception:
            n_rows = -1
        sheet = str(row["sheet_name"])
        if sheet not in xls.sheet_names or n_rows <= 0:
            n_rows_ok = False
            n_rows_details.append(f"{sheet}: invalid n_rows={row.get('n_rows')}")
            continue

        try:
            data = _read_measurement_data(workbook, sheet)
        except Exception as exc:
            n_rows_ok = False
            n_rows_details.append(f"{sheet}: {exc}")
            continue

        actual_rows = len(data.dropna(how="all"))
        if actual_rows != n_rows:
            n_rows_ok = False
        n_rows_details.append(f"{sheet}: index={n_rows}, data={actual_rows}")

    checks.append({
        "check": "data_row_counts",
        "status": "PASS" if n_rows_ok else "FAIL",
        "detail": "; ".join(n_rows_details),
    })

    xy_ok = True
    xy_details: list[str] = []
    for _, row in index.iterrows():
        sheet = str(row["sheet_name"])
        source_text = str(row.get("source_file", "")).strip()
        source_path = Path(source_text)

        if not source_text or not source_path.exists():
            xy_ok = False
            xy_details.append(f"{sheet}: source missing ({source_text})")
            continue

        try:
            exported = _read_measurement_data(workbook, sheet)
            source = pd.read_csv(source_path)
        except Exception as exc:
            xy_ok = False
            xy_details.append(f"{sheet}: read error ({exc})")
            continue

        required = ["distance_m", "depth_m"]
        if any(col not in exported.columns for col in required) or any(
            col not in source.columns for col in required
        ):
            xy_ok = False
            xy_details.append(f"{sheet}: distance_m/depth_m missing")
            continue

        exp_xy = exported[required].dropna(how="all").reset_index(drop=True)
        src_xy = source[required].dropna(how="all").reset_index(drop=True)

        if len(exp_xy) != len(src_xy):
            xy_ok = False
            xy_details.append(
                f"{sheet}: row mismatch exported={len(exp_xy)}, source={len(src_xy)}"
            )
            continue

        exp_values = exp_xy.apply(pd.to_numeric, errors="coerce").to_numpy(dtype=float)
        src_values = src_xy.apply(pd.to_numeric, errors="coerce").to_numpy(dtype=float)
        matched = np.allclose(
            exp_values,
            src_values,
            rtol=1e-10,
            atol=1e-12,
            equal_nan=True,
        )
        if not matched:
            xy_ok = False
        xy_details.append(
            f"{sheet}: rows={len(exp_xy)}, distance/depth={'match' if matched else 'DIFFER'}"
        )

    checks.append({
        "check": "xy_source_fidelity",
        "status": "PASS" if xy_ok else "FAIL",
        "detail": "; ".join(xy_details),
    })

    wb = load_workbook(workbook, read_only=False, data_only=False)
    chart_ok = True
    chart_details: list[str] = []
    for sheet in index["sheet_name"]:
        if sheet not in wb.sheetnames:
            chart_ok = False
            chart_details.append(f"{sheet}: missing")
            continue
        n_charts = len(wb[sheet]._charts)
        if n_charts != 1:
            chart_ok = False
        chart_details.append(f"{sheet}: charts={n_charts}")

    checks.append({
        "check": "one_chart_per_measurement",
        "status": "PASS" if chart_ok else "FAIL",
        "detail": "; ".join(chart_details),
    })

    report = pd.DataFrame(checks)
    report_dir = Path("runs_acceptance/_checks")
    report_dir.mkdir(parents=True, exist_ok=True)
    report_path = report_dir / "section_profiles_acceptance.csv"
    report.to_csv(report_path, index=False, encoding="utf-8-sig")

    n_pass = int(report["status"].eq("PASS").sum())
    n_fail = int(report["status"].eq("FAIL").sum())

    print("Aforix section profiles acceptance")
    print("=================================")
    print(f"Workbook: {workbook}")
    print(f"Checks: {len(report)}")
    print(f"PASS: {n_pass}")
    print(f"FAIL: {n_fail}")
    print(f"Report: {report_path}")
    print("")
    print(report.to_string(index=False))

    if n_fail:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
