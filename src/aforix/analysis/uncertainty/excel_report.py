"""Auditable XLSX report for a single IVE calculation."""
from __future__ import annotations

from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from aforix.analysis.uncertainty.calculator import IVECalculation, IVEParameters


def _safe_name(value: object) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(value)).strip("_") or "unknown"


def write_ive_excel_report(
    output_root: str | Path,
    *,
    instrument: str,
    station_id: object,
    measurement_date: object,
    measurement_time: object,
    normalized_points,
    result: IVECalculation,
    parameters: IVEParameters | None = None,
) -> Path:
    """Write a traceable, values-only workbook; never re-evaluate the IVE model."""
    p = parameters or IVEParameters()
    root = Path(output_root)
    root.mkdir(parents=True, exist_ok=True)
    filename = "IVE_{}_{}_{}_{}.xlsx".format(
        *(_safe_name(v) for v in (instrument, station_id, measurement_date, measurement_time))
    )
    target = root / filename

    wb = Workbook()
    summary = wb.active
    summary.title = "Resumen"
    fields = [
        ("Instrumento", instrument),
        ("Estacion", str(station_id)),
        ("Fecha", str(measurement_date)),
        ("Hora", str(measurement_time)),
        ("Estado IVE", result.status),
        ("Calidad IVE", result.quality),
        ("Verticales medidas", result.vertical_count),
        ("Caudal Mid-section (m3/s)", result.q_mid_m3s),
        ("Incertidumbre relativa (%)", result.relative_uncertainty_pct),
        ("Incertidumbre absoluta (m3/s)", result.absolute_uncertainty_m3s),
        ("s_d IVE (m)", result.sd_depth_ive_m),
        ("s_v IVE (m/s)", result.sd_velocity_ive_m_s),
        ("s_d utilizado (m)", result.sd_depth_used_m),
        ("s_v utilizado (m/s)", result.sd_velocity_used_m_s),
        ("Componente ancho (relativa)", result.uncertainty_width_relative),
        ("Componente profundidad (relativa)", result.uncertainty_depth_relative),
        ("Componente velocidad (relativa)", result.uncertainty_velocity_relative),
        ("Componente sistematica (relativa)", p.systematic_relative),
    ]
    summary.append(["Parametro", "Valor"])
    for key, value in fields:
        summary.append([key, value])
    summary.append(["Nota", "Resultados copiados del motor IVE; sin formulas independientes."])
    if result.status == "provisional":
        summary.append(["Advertencia", "Resultado provisional: falta minimo instrumental de velocidad."])

    settings = wb.create_sheet("Parametros")
    settings.append(["Parametro", "Valor"])
    for key, value in vars(p).items():
        settings.append([key, value])
    settings.append(["Tratamiento bordes", "Nodos fisicos con profundidad y velocidad nulas cuando faltan"])
    settings.append(["Metodo", "Mid-section; se conservan las velocidades medidas extremas"])

    source = wb.create_sheet("Points_normalizados")
    columns = list(normalized_points.columns)
    source.append(columns)
    for row in normalized_points.itertuples(index=False, name=None):
        source.append([v.item() if hasattr(v, "item") else (None if v is None else v) for v in row])

    details = wb.create_sheet("Detalle_IVE")
    keys = [
        "index", "distance_m", "depth_m", "velocity_m_s", "width_mid_m",
        "q_mid_m3s", "interior", "ive_residual", "omega",
        "depth_estimated_m", "velocity_estimated_m_s",
        "depth_residual", "velocity_residual", "variance_width",
        "variance_depth", "variance_velocity",
    ]
    details.append(keys)
    for row in result.details:
        details.append([row.get(key) for key in keys])

    for sheet in wb.worksheets:
        sheet.freeze_panes = "B2" if sheet.title in ("Resumen", "Parametros") else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        for col in sheet.columns:
            letter = get_column_letter(col[0].column)
            max_length = max(len(str(cell.value or "")) for cell in col)
            sheet.column_dimensions[letter].width = min(max(14, max_length + 2), 65)

    wb.save(target)
    return target
