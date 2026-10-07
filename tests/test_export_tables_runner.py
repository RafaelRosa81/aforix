from __future__ import annotations

from openpyxl import load_workbook
import pandas as pd
import pytest

from aforix.export.tables.runner import ExportRequest, run_export_tables


def test_flat_xlsx_preserves_leading_zero_measurement_time(tmp_path):
    normalized_root = tmp_path / "database" / "normalized"
    normalized_root.mkdir(parents=True)

    pd.DataFrame(
        [
            {
                "instrument": "flowtracker",
                "station_id": "7005",
                "station_name": "COLORADO_R36",
                "measurement_date": "20251217",
                "measurement_time": "093425",
                "q_total_m3s": 0.0498,
            }
        ]
    ).to_csv(normalized_root / "Summary.csv", index=False)

    config = {
        "__repo_root__": str(tmp_path),
        "export": {
            "tables": {
                "input_dir": "database/normalized",
                "output_dir": "outputs/tables",
            }
        },
    }

    result = run_export_tables(
        config,
        ExportRequest(
            table="Summary",
            instrument="all",
            parameters=("q_total_m3s",),
            grouping="none",
            fmt="xlsx",
            pivot=False,
        ),
    )

    wb = load_workbook(result.output_file, data_only=False)
    ws = wb["export"]

    headers = [cell.value for cell in ws[1]]
    time_col = headers.index("measurement_time") + 1
    station_col = headers.index("station_id") + 1
    date_col = headers.index("measurement_date") + 1

    assert ws.cell(2, station_col).value == "7005"
    assert ws.cell(2, date_col).value == "20251217"
    assert ws.cell(2, time_col).value == "093425"
    assert ws.cell(2, time_col).data_type == "s"


@pytest.mark.parametrize("point_column", ["Point", "point", "station", "site_id"])
def test_flat_export_preserves_leading_zero_point_aliases(tmp_path, point_column):
    normalized_root = tmp_path / "database" / "normalized"
    normalized_root.mkdir(parents=True)
    (normalized_root / "Summary.csv").write_text(
        f"{point_column},q_total_m3s\n0012,1.0\n12,2.0\n",
        encoding="utf-8",
    )

    config = {
        "__repo_root__": str(tmp_path),
        "export": {
            "tables": {
                "input_dir": "database/normalized",
                "output_dir": "outputs/tables",
            }
        },
    }

    result = run_export_tables(
        config,
        ExportRequest(
            table="Summary",
            points=("0012",),
            parameters=("q_total_m3s",),
            grouping="none",
            fmt="csv",
            pivot=False,
        ),
    )

    exported = pd.read_csv(result.output_file, dtype={point_column: "string"})
    assert exported[point_column].tolist() == ["0012"]
    assert exported["q_total_m3s"].tolist() == [1.0]

