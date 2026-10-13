from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from aforix.analysis.uncertainty.calculator import IVEPoint, calculate_ive
from aforix.analysis.uncertainty.excel_report import write_ive_excel_report


def test_ive_report_preserves_engine_values_and_normalized_inputs(tmp_path: Path):
    x = [0.2 + 0.4 * i for i in range(10)]
    depths = [.1862, .2348, .1959, .177, .1768, .1671, .1623, .1697, .1989, .1967]
    velocities = [.0463, .1218, .0944, .077, .1593, .1337, .1349, .1362, .0582, .0906]
    df = pd.DataFrame({
        "distance_m": x,
        "depth_m": depths,
        "velocity_mean_m_s": velocities,
        "width_m": [4.0] * 10,
    })
    result = calculate_ive(
        [IVEPoint(a, b, c) for a, b, c in zip(x, depths, velocities)],
        width_m=4.0,
    )
    path = write_ive_excel_report(
        tmp_path,
        instrument="nivus",
        station_id="7001",
        measurement_date="20241219",
        measurement_time="214313",
        normalized_points=df,
        result=result,
    )
    assert path.is_file()
    wb = load_workbook(path, data_only=True)
    summary = dict(wb["Resumen"].values)
    assert summary["Caudal Mid-section (m3/s)"] == result.q_mid_m3s
    assert summary["Incertidumbre relativa (%)"] == result.relative_uncertainty_pct
    assert summary["Incertidumbre absoluta (m3/s)"] == result.absolute_uncertainty_m3s
    assert summary["Estado IVE"] == "calculated"
    assert "no se aplica s_v,min" in summary["Decision metodologica"]
    assert wb["Points_normalizados"]["A2"].value == x[0]
    assert wb["Detalle_IVE"].max_row == len(result.details) + 1
    assert len(result.details) == 12
